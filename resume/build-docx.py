#!/usr/bin/env python3
"""Writes public/assets/omid-heidari-resume.docx from omid-heidari-resume.html. Edit the HTML, not this file."""

import re
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / 'omid-heidari-resume.html'
OUT = HERE.parent / 'public' / 'assets' / 'omid-heidari-resume.docx'

# Arial and Consolas, not Avenir Next and DM Mono: a .docx is re-rendered with the reader's fonts, and Windows and
# Microsoft Office ship only the first pair.
FONT, MONO = 'Arial', 'Consolas'
INK, DIM, FAINT, ACCENT, RULE = (RGBColor.from_string(c) for c in ('141414', '474747', '6B6B6B', 'C74D14', 'D6D6D6'))
TEXT_WIDTH = Mm(210 - 14 - 14)
BULLET_NUM_ID = 90


class Node:
	def __init__(self, tag, attrs, parent):
		self.tag, self.attrs, self.parent, self.children = tag, dict(attrs), parent, []

	@property
	def classes(self):
		return set((self.attrs.get('class') or '').split())

	def elements(self):
		return [c for c in self.children if isinstance(c, Node)]

	def find(self, tag):
		for child in self.elements():
			if child.tag == tag:
				return child
			hit = child.find(tag)
			if hit:
				return hit
		return None

	def text(self):
		return ''.join(c if isinstance(c, str) else c.text() for c in self.children)


class TreeBuilder(HTMLParser):
	VOID = {'meta', 'br', 'img', 'link', 'hr', 'input'}

	def __init__(self):
		super().__init__(convert_charrefs=True)
		self.root = self.cur = Node('#root', [], None)

	def handle_starttag(self, tag, attrs):
		node = Node(tag, attrs, self.cur)
		self.cur.children.append(node)
		if tag not in self.VOID:
			self.cur = node

	def handle_endtag(self, tag):
		node = self.cur
		while node.parent and node.tag != tag:
			node = node.parent
		if node.parent:
			self.cur = node.parent

	def handle_data(self, data):
		self.cur.children.append(data)


def runs_of(el, fmt):
	if 'sep' in el.classes:
		return [[' | ', {**fmt, 'color': FAINT}]]
	f = dict(fmt)
	if el.tag in ('h3', 'h4') or el.classes & {'label', 'accent'}:
		f['bold'] = True
	if 'accent' in el.classes:
		f['color'] = ACCENT
	if 'last' in el.classes:
		f.update(bold=False, color=FAINT)
	if 'meta' in el.classes:
		f.update(bold=False, color=DIM)
	if el.classes & {'when', 'links', 'stack'}:
		f.update(font=MONO, size=8, color=INK if 'stack' in el.classes else DIM)
	if 'links' in el.classes:
		f['link_color'] = ACCENT
	if el.tag == 'a':
		f.update(link=el.attrs['href'], color=f.get('link_color', INK))
	out = []
	for child in el.children:
		out += [[re.sub(r'[ \t\n\r\f\v]+', ' ', child), f]] if isinstance(child, str) else runs_of(child, f)
	return out


def tidy(runs):
	out = []
	for text, fmt in runs:
		if out and (text == '\t' or out[-1][0] == '\t'):
			if text == '\t':
				out[-1][0] = out[-1][0].rstrip()
			else:
				text = text.lstrip()
		elif out and out[-1][0].endswith(' ') and text.startswith(' '):
			text = text[1:]
		if text:
			out.append([text, fmt])
	if out:
		out[0][0] = out[0][0].lstrip()
		out[-1][0] = out[-1][0].rstrip(' ')
	return [r for r in out if r[0]]


def set_fonts(rpr, font=FONT):
	fonts = rpr.find(qn('w:rFonts'))
	if fonts is None:
		fonts = OxmlElement('w:rFonts')
		rpr.insert(0, fonts)
	for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
		fonts.attrib.pop(qn(f'w:{attr}'), None)
	for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
		fonts.set(qn(f'w:{attr}'), font)


def setup(doc, title):
	section = doc.sections[0]
	section.page_width, section.page_height = Mm(210), Mm(297)
	section.top_margin, section.bottom_margin = Mm(12), Mm(11)
	section.left_margin = section.right_margin = Mm(14)

	set_fonts(doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr')))

	normal = doc.styles['Normal']
	normal.font.name = FONT
	normal.font.size = Pt(9)
	normal.font.color.rgb = INK
	set_fonts(normal.element.get_or_add_rPr())
	pf = normal.paragraph_format
	pf.space_before = pf.space_after = Pt(0)
	pf.line_spacing = 1.2
	pf.widow_control = True

	name = doc.styles.add_style('Resume Name', WD_STYLE_TYPE.PARAGRAPH)
	name.base_style = normal
	name.font.size = Pt(24)
	name.font.bold = True
	name.paragraph_format.line_spacing = 1.0

	heading = doc.styles.add_style('Resume Section', WD_STYLE_TYPE.PARAGRAPH)
	heading.base_style = normal
	heading.font.size = Pt(11.5)
	heading.font.bold = True
	heading.font.color.rgb = ACCENT
	hpf = heading.paragraph_format
	hpf.space_before = Pt(9)
	hpf.space_after = Pt(3)
	hpf.line_spacing = 1.1
	hpf.keep_with_next = True
	hpf.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
	heading.element.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))

	numbering = doc.part.numbering_part.element
	abstract = parse_xml(
		f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{BULLET_NUM_ID}">'
		'<w:multiLevelType w:val="singleLevel"/>'
		'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
		'<w:pPr><w:ind w:left="227" w:hanging="170"/></w:pPr>'
		f'<w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}"/><w:color w:val="{ACCENT}"/></w:rPr></w:lvl></w:abstractNum>'
	)
	num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{BULLET_NUM_ID}"><w:abstractNumId w:val="{BULLET_NUM_ID}"/></w:num>')
	nums = numbering.findall(qn('w:num'))
	if nums:
		nums[0].addprevious(abstract)
		nums[-1].addnext(num)
	else:
		numbering.append(abstract)
		numbering.append(num)

	for setting in doc.settings.element.iter(qn('w:compatSetting')):
		if setting.get(qn('w:name')) == 'compatibilityMode':
			setting.set(qn('w:val'), '15')
	zoom = doc.settings.element.find(qn('w:zoom'))
	if zoom is not None and zoom.get(qn('w:percent')) is None:
		zoom.set(qn('w:percent'), '100')

	props = doc.core_properties
	props.title = title
	props.author = props.last_modified_by = 'Omid Heidari'
	props.subject = 'Resume'
	props.keywords = 'Front-End Developer, React, Next.js, TypeScript'
	props.comments = ''
	props.revision = 1
	props.created = props.modified = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)


class Builder:
	def __init__(self, doc):
		self.doc = doc
		self.gap = 0

	def par(self, runs, *, style=None, after=0, tab=False, bullet=False, keep=False):
		p = self.doc.add_paragraph(style=style)
		pf = p.paragraph_format
		if self.gap:
			pf.space_before = Pt(self.gap)
			self.gap = 0
		if after:
			pf.space_after = Pt(after)
		if keep:
			pf.keep_with_next = True
		if tab:
			pf.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
		if bullet:
			num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
			num_pr.get_or_add_ilvl().val = 0
			num_pr.get_or_add_numId().val = BULLET_NUM_ID
		for text, fmt in tidy(runs):
			run = p.add_run(text)
			if 'bold' in fmt:
				run.font.bold = fmt['bold']
			if 'font' in fmt:
				set_fonts(run._r.get_or_add_rPr(), fmt['font'])
			if 'size' in fmt:
				run.font.size = Pt(fmt['size'])
			if fmt.get('color') is not None:
				run.font.color.rgb = fmt['color']
			if 'raise' in fmt:
				run._r.get_or_add_rPr().find(qn('w:sz')).addprevious(
					parse_xml(f'<w:position {nsdecls("w")} w:val="{round(fmt["raise"] * 2)}"/>')
				)
			if fmt.get('link'):
				run.font.underline = False
				link = OxmlElement('w:hyperlink')
				link.set(qn('r:id'), p.part.relate_to(fmt['link'], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
				run._r.addprevious(link)
				link.append(run._r)
		return p

	def header(self, el):
		for child in el.elements():
			if child.tag == 'h1':
				self.par(runs_of(child, {}), style='Resume Name')
			elif 'headline' in child.classes:
				self.gap = 2
				self.par(runs_of(child, {'size': 10.5}))
			elif 'contact' in child.classes:
				self.gap = 3
				self.par(runs_of(child, {'size': 8.5, 'color': DIM}))

	def block(self, el, prev=None):
		c = el.classes
		if el.tag == 'h2':
			# The tab's leader line is the PDF's rule beside the heading, raised from the baseline to mid-height. The gap
			# before it is ASCII spaces, which every parser trims, so the heading still reads as one word.
			self.par(runs_of(el, {}) + [['   \t', {'bold': False, 'size': 6, 'color': RULE, 'raise': 3.5}]], style='Resume Section')
		elif el.tag == 'section' or c & {'entry', 'sub'}:
			if 'entry' in c and prev is not None and 'entry' in prev.classes:
				self.gap = 7
			if 'sub' in c:
				self.gap = 4
			children = el.elements()
			for i, child in enumerate(children):
				self.block(child, children[i - 1] if i else None)
		elif 'row' in c:
			left, right = el.elements()
			if prev is not None and 'row' in prev.classes:
				self.gap = 1.5
			self.par(runs_of(left, {'size': 9.5} if left.tag == 'h3' else {}) + [['\t', {}]] + runs_of(right, {}), tab=True, keep=True)
		elif el.tag == 'ul':
			self.gap = max(self.gap, 1)
			for li in el.elements():
				self.par(runs_of(li, {}), bullet=True, after=0.8)
		elif el.tag == 'p':
			if 'lead' in c:
				self.gap = 0.5
			if 'stack' in c:
				self.gap = 0.3
			if 'training' in c:
				self.gap = 3
			if prev is not None and prev.tag == 'p' and not prev.classes:
				self.gap = 1.5
			self.par(runs_of(el, {'color': DIM} if c & {'lead', 'training'} else {}), keep='stack' in c)


def main():
	tree = TreeBuilder()
	tree.feed(SRC.read_text(encoding='utf-8'))
	body = tree.root.find('body')
	doc = Document()
	setup(doc, re.sub(r'\s+', ' ', tree.root.find('title').text()).strip())
	builder = Builder(doc)
	for el in body.elements():
		if el.tag == 'header':
			builder.header(el)
		elif el.tag == 'section':
			builder.block(el)
	doc.save(OUT)
	print(f'Wrote {OUT}')


if __name__ == '__main__':
	main()
